import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from docx import Document
import io

load_dotenv()
client = OpenAI()

st.title("EXERCISE 3")

## STEP 1: 创建文本框，接收用户输入的 hypothetical scenario
scenario = st.text_area(
    "Hypothetical Scenario Involving Contract Formation",
    height=200,
    placeholder="e.g., describe the parties involved, what was offered, and how it was accepted..."
)

## 合同构成的六个要素（名称 + 定义），按顺序逐一分析
elements = [
    ("Offer and Acceptance",
     "A contract requires a definite offer by one party that is accepted, exactly as proposed, by the other party. A mere willingness to negotiate is not an offer. Acceptance must be a clear and unequivocal response communicated to the offeror, and can be shown through words or conduct."),

    ("Intention to Create Legal Relations",
     "The parties must have intended their agreement to be legally binding, assessed objectively. Commercial/arm's-length dealings are presumed to carry this intention; social or domestic arrangements are presumed NOT to, unless stated otherwise."),

    ("Consideration",
     "Each party must give something of value in exchange for the other's promise (not necessarily money). Love, affection, or a pure gift do not count as valid consideration."),

    ("Legal Capacity",
     "Both parties must have the legal capacity to contract. Minors, people with mental impairment, bankrupts, corporations acting without authority, and prisoners may have limited or conditional capacity."),

    ("Consent",
     "Each party's agreement must be genuine and freely given, not undermined by mistake, misrepresentation, duress, undue influence, or unfair terms."),

    ("Legality",
     "The contract must not involve illegal conduct or be contrary to public policy.")
]

## 用 session_state 记录当前进行到第几个要素、已分析结果、是否提前终止
if "current_step" not in st.session_state:
    st.session_state.current_step = 0
    st.session_state.results = []
    st.session_state.stopped = False

## 封装函数：只分析"当前这一个"要素，分析完立刻强制重跑页面
def analyse_current_element():
    name, definition = elements[st.session_state.current_step]
    prompt = f"""
    You are a contract law assistant analysing ONLY the element below.

    Element: {name}
    Definition: {definition}

    Scenario: {scenario}

    Respond in EXACTLY this format, nothing else:
    VERDICT: YES or NO
    REASON: one or two sentence explanation referring to the facts.
    """

    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "user", "content": prompt}]
    )
    answer = response.choices[0].message.content.strip()
    st.session_state.results.append((name, answer))

    first_line = answer.splitlines()[0].upper()
    if "NO" in first_line:
        st.session_state.stopped = True
    else:
        st.session_state.current_step += 1

    st.rerun()   ## 强制立刻重新跑一遍脚本，确保按钮更新到下一个要素

## STEP 2 & 3: 按钮触发"下一步"分析，结果逐个显示在屏幕上
if scenario and not st.session_state.stopped and st.session_state.current_step < len(elements):
    step_name = elements[st.session_state.current_step][0]
    if st.button(f"Analyse: {step_name}"):
        analyse_current_element()

## 显示已经分析过的每一个要素（按顺序，逐个出现）
for name, answer in st.session_state.results:
    st.subheader(name)
    st.write(answer)

## 最终结论
if st.session_state.stopped:
    st.error("❌ Overall conclusion: A contract has NOT been formed.")

if (not st.session_state.stopped) and st.session_state.current_step == len(elements) and len(elements) > 0:
    st.success("✅ Overall conclusion: A contract has been formed.")

## STEP 4: 追问功能（依赖已经至少分析过一个要素）
if len(st.session_state.results) > 0:
    followup = st.text_input("Ask a follow-up question about this analysis")

    if st.button("Ask Follow-up"):
        previous_analysis = "\n\n".join(
            [f"{name}:\n{answer}" for name, answer in st.session_state.results]
        )

        followup_prompt = f"""
        Here is your previous element-by-element analysis of a contract formation scenario:

        {previous_analysis}

        The user has a follow-up question about this analysis:
        {followup}

        Answer the follow-up question, staying consistent with your previous analysis.
        """

        followup_response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": followup_prompt}]
        )
        st.session_state.followup_response = followup_response.choices[0].message.content

    if "followup_response" in st.session_state:
        st.write(st.session_state.followup_response)

## STEP 5: 导出为 .docx 文件
if len(st.session_state.results) > 0:
    doc = Document()
    doc.add_heading("Contract Formation Analysis", level=1)

    for name, answer in st.session_state.results:
        doc.add_heading(name, level=2)
        doc.add_paragraph(answer)

    if st.session_state.stopped:
        conclusion = "A contract has NOT been formed."
    elif st.session_state.current_step == len(elements):
        conclusion = "A contract has been formed."
    else:
        conclusion = "Analysis in progress."

    doc.add_heading("Overall Conclusion", level=2)
    doc.add_paragraph(conclusion)

    if "followup_response" in st.session_state:
        doc.add_heading("Follow-up Answer", level=2)
        doc.add_paragraph(st.session_state.followup_response)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.download_button(
        label="Download as Word Document",
        data=buffer,
        file_name="contract_analysis.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )